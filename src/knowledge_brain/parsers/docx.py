from io import BytesIO

import docx

from knowledge_brain.parsers.base import ParseError


def parse_docx(content: bytes) -> str:
    try:
        document = docx.Document(BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text for c in row.cells if c.text]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"failed to parse DOCX: {exc}") from exc
