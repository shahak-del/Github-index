from io import BytesIO

from pptx import Presentation

from knowledge_brain.parsers.base import ParseError


def parse_pptx(content: bytes) -> str:
    try:
        presentation = Presentation(BytesIO(content))
        parts = []
        for i, slide in enumerate(presentation.slides, start=1):
            slide_parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text
                    if text:
                        slide_parts.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text for c in row.cells if c.text]
                        if cells:
                            slide_parts.append(" | ".join(cells))
            if slide_parts:
                parts.append(f"# Slide {i}\n" + "\n".join(slide_parts))
        return "\n\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"failed to parse PPTX: {exc}") from exc
